"""Generadores de PDF, DOCX y XLSX sintéticos para las pruebas (nada binario se versiona)."""

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

import docx
import openpyxl
import pymupdf
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml

# Caso de aceptación del SKILL: 27 códigos generales y 9 apariciones en Procedimientos.
ACCEPTANCE_GENERAL = [
    "PO-SGS-400", "PO-SGS-615", "ES-SGC-001", "AL-SGS-401", "AL-SGSA-911", "PC-SGSA-914",
    "LC-SGSA-915", "ES-SGSA-919", "AX-SGC-002", "ES-SGS-404", "ES-SGS-407", "ES-SGS-416",
    "ES-SGS-525", "ES-SGS-546", "ES-SGG-547", "ES-SGG-597", "ES-COM-001", "ES-SGC-002",
    "ES-SGC-100", "ES-AFP-001", "ES-SOP-001", "LC-SGG-050", "ES-SGC-210", "AX-SGC-001",
    "FC-SGC-100", "LC-SGC-500", "LC-SGS-849",
]  # fmt: skip
MANUAL_HEADER = "MC-SGG-435 / 23 - Manual de los Sistemas de Gestión - Página {n}/{total}"
ACCEPTANCE_PROCEDURES = [
    "PR-SGG-000", "PR-SGG-101", "PR-SGG-004", "PR-SGG-005", "PR-SGC-006", "PR-SGG-007",
    "PR-SGG-101", "PR-SGG-102", "PR-SGS-402",
]  # fmt: skip


# --------------------------------------------------------------------------- PDF


@dataclass(frozen=True)
class Line:
    text: str
    bold: bool = False
    size: float = 10.0
    x: float = 72.0
    # Texto adicional en la misma línea base (simula una segunda columna o celda).
    right: str | None = None


def heading(text: str) -> Line:
    return Line(text, bold=True, size=13.0)


def ref(code: str, title: str = "Documento referenciado") -> Line:
    return Line(f"[{code}]", right=title)


def build_pdf(
    path: Path,
    pages: list[list[Line]],
    *,
    header: str | None = None,
    footer: str | None = "Pza. Independencia 822 Piso 4",
    side_mark: str | None = "CONFIDENCIAL TILSOR",
    bookmarks: list[list[int | str]] | None = None,
    scanned_pages: set[int] = frozenset(),
    encrypt: bool = False,
) -> Path:
    """Crea un PDF con texto real; las páginas en ``scanned_pages`` son solo una imagen."""
    document = pymupdf.open()
    total = len(pages)
    for number, lines in enumerate(pages, start=1):
        page = document.new_page(width=595, height=842)
        if number in scanned_pages:
            pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), False)
            pixmap.clear_with(220)
            page.insert_image(page.rect, pixmap=pixmap)
            continue
        if header:
            page.insert_text((300, 60), header.format(n=number, total=total), fontsize=8)
        if footer:
            page.insert_text((72, 800), footer, fontsize=8)
        if side_mark:
            page.insert_text((30, 600), side_mark, fontsize=12, rotate=90)
        y = 180.0
        for line in lines:
            font = "hebo" if line.bold else "helv"
            page.insert_text((line.x, y), line.text, fontname=font, fontsize=line.size)
            if line.right:
                page.insert_text((200, y), line.right, fontname="helv", fontsize=line.size)
            y += line.size * 1.7
    if bookmarks:
        document.set_toc(bookmarks)
    options = {}
    if encrypt:
        options = {
            "encryption": pymupdf.PDF_ENCRYPT_AES_256,
            "owner_pw": "propietario",
            "user_pw": "usuario",
        }
    document.save(path, **options)
    document.close()
    return path


def manual_pages() -> list[list[Line]]:
    """Réplica sintética del manual MC-SGG-435 con todas sus trampas conocidas."""
    cover = [
        Line("MC-SGG-435 - Manual de los Sistemas de Gestión", bold=True, size=18),
        Line("Tipo de documento : Manual"),
        Line("Referencia/Versión : MC-SGG-435/23"),
        Line("Clasificación : CONFIDENCIAL TILSOR"),
    ]
    toc = [
        Line("Tabla de Contenido", bold=True, size=15),
        Line("1 OBJETIVO ........................................ 3", bold=True),
        Line("3 RESPONSABILIDADES ............................... 3", bold=True),
        Line("4 REFERENCIAS ..................................... 3", bold=True),
        Line("5 HISTORIAL ....................................... 4", bold=True),
    ]
    body = [
        heading("1 OBJETIVO"),
        Line("Describir los Sistemas de Gestión conforme a las normas ISO."),
        heading("3 RESPONSABILIDADES"),
        Line("Los responsables se describen en su perfil de cargo en:"),
        Line("SGSI", right="[ES-SGG-547] Perfiles de Cargo Comité de Seguridad"),
        Line("SGC y SGSA", right="[ES-SGC-100] Perfiles de Cargo Dirección de Tilsor"),
        heading("4 REFERENCIAS"),
        *(ref(code) for code in ACCEPTANCE_GENERAL[:14]),
    ]
    continuation = [
        *(ref(code) for code in ACCEPTANCE_GENERAL[14:]),
        Line("Procedimientos", bold=True),
        *(ref(code) for code in ACCEPTANCE_PROCEDURES),
        heading("5 HISTORIAL"),
        Line("23 06/07/2026 Actualización Se cambia referencia de ES-SGS-547 por ES-SGG-547"),
        heading("6 DEFINICIONES"),
        Line("Consultar el Glosario de los Sistemas de Gestión [ES-SGG-597]."),
        heading("7 DESCRIPCIÓN"),
        Line("Los niveles de servicio se detallan en [AC-AFP-698]."),
    ]
    return [cover, toc, body, continuation]


# --------------------------------------------------------------------------- DOCX


@dataclass(frozen=True)
class Heading:
    text: str
    level: int = 1


@dataclass(frozen=True)
class Para:
    text: str
    bold: bool = False


@dataclass(frozen=True)
class TocEntry:
    text: str


@dataclass(frozen=True)
class TableBlock:
    rows: list[list[str]]


@dataclass(frozen=True)
class TextBox:
    text: str


DocxItem = Heading | Para | TocEntry | TableBlock | TextBox

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def build_docx(
    path: Path,
    items: list[DocxItem],
    *,
    localized_heading_ids: bool = False,
    custom_properties: dict[str, str] | None = None,
    header_text: str | None = None,
) -> Path:
    document = docx.Document()
    if localized_heading_ids:
        # Como en Word en español: el ID del estilo cambia, el nombre interno no.
        document.styles["Heading 1"].style_id = "Ttulo1"
        document.styles["Heading 2"].style_id = "Ttulo2"
    if header_text:
        document.sections[0].header.paragraphs[0].text = header_text
    for item in items:
        if isinstance(item, Heading):
            document.add_heading(item.text, level=item.level)
        elif isinstance(item, Para):
            paragraph = document.add_paragraph()
            paragraph.add_run(item.text).bold = item.bold or None
        elif isinstance(item, TocEntry):
            if "toc 1" not in {s.name.casefold() for s in document.styles}:
                document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
            document.add_paragraph(item.text, style="toc 1")
        elif isinstance(item, TableBlock):
            table = document.add_table(rows=len(item.rows), cols=len(item.rows[0]))
            for r, row in enumerate(item.rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        elif isinstance(item, TextBox):
            document.add_paragraph()._p.append(_textbox_run(item.text))
    document.save(str(path))
    if custom_properties:
        _add_custom_properties(path, custom_properties)
    return path


def _textbox_run(text: str) -> object:
    """Cuadro de texto con la copia Choice/Fallback que genera Word."""
    content = f"<w:txbxContent><w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p></w:txbxContent>"
    return parse_xml(
        f'<w:r xmlns:w="{_W}" xmlns:mc="{_MC}"><mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:drawing>{content}</w:drawing></mc:Choice>'
        f"<mc:Fallback><w:pict>{content}</w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )


def _add_custom_properties(path: Path, properties: dict[str, str]) -> None:
    fmtid = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"
    props = "".join(
        f'<property fmtid="{fmtid}" pid="{pid}" name="{escape(name)}">'
        f"<vt:lpwstr>{escape(value)}</vt:lpwstr></property>"
        for pid, (name, value) in enumerate(properties.items(), start=2)
    )
    custom = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"'
        ' xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f"{props}</Properties>"
    )
    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries["docProps/custom.xml"] = custom.encode()
    entries["[Content_Types].xml"] = re.sub(
        rb"</Types>",
        b'<Override PartName="/docProps/custom.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
        entries["[Content_Types].xml"],
    )
    entries["_rels/.rels"] = re.sub(
        rb"</Relationships>",
        b'<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/'
        b'2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>',
        entries["_rels/.rels"],
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries.items():
            target.writestr(name, data)


# --------------------------------------------------------------------------- XLSX


def build_xlsx(
    path: Path,
    *,
    title: str | None = None,
    custom_properties: dict[str, str] | None = None,
    rows: list[list[str]] | None = None,
) -> Path:
    """Planilla mínima; el contenido no se analiza (solo identidad y propiedades)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    if title:
        workbook.properties.title = title
    for row in rows or [["Código", "Descripción"], ["X-1", "[PO-SGS-400] no se analiza"]]:
        workbook.active.append(row)
    workbook.save(path)
    if custom_properties:
        _add_custom_properties(path, custom_properties)
    return path


def manual_docx_items(*, references_in_table: bool = False) -> list[DocxItem]:
    """El caso de aceptación en Word: numeración en el estilo, no en el texto."""
    references: list[DocxItem]
    if references_in_table:
        rows = [[f"[{code}]", "Documento"] for code in ACCEPTANCE_GENERAL]
        rows.append(["Procedimientos", ""])
        rows += [[f"[{code}]", "Procedimiento"] for code in ACCEPTANCE_PROCEDURES]
        references = [TableBlock(rows)]
    else:
        references = [
            *(Para(f"[{code}]\tDocumento") for code in ACCEPTANCE_GENERAL),
            Para("Procedimientos", bold=True),
            *(Para(f"[{code}]\tProcedimiento") for code in ACCEPTANCE_PROCEDURES),
        ]
    return [
        Para("Referencia/Versión : MC-SGG-435/23"),
        TocEntry("REFERENCIAS\t3"),
        TocEntry("HISTORIAL\t4"),
        Heading("OBJETIVO"),
        Para("Describir los Sistemas de Gestión."),
        Heading("RESPONSABILIDADES"),
        TableBlock([["SGSI", "[ES-SGG-547]", "Perfiles de Cargo Comité de Seguridad"]]),
        Heading("REFERENCIAS"),
        *references,
        Heading("HISTORIAL"),
        Para("Se cambia referencia de ES-SGS-547 por ES-SGG-547"),
        Heading("DEFINICIONES"),
        Para("Consultar el Glosario [ES-SGG-597]."),
    ]
