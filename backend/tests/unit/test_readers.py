from pathlib import Path

import pytest

from app.ingestion.readers import EncryptedDocumentError, ReadError, read_document
from app.ingestion.readers.docx import read_docx
from app.ingestion.readers.pdf import read_pdf
from app.ingestion.readers.xlsx import read_xlsx
from app.models import FileFormat
from tests.fixtures.builders import (
    MANUAL_HEADER,
    Heading,
    Line,
    Para,
    TableBlock,
    TextBox,
    TocEntry,
    build_docx,
    build_pdf,
    build_xlsx,
    manual_pages,
)


class TestPdfReader:
    def test_lines_keep_page_order_and_join_cells_on_the_same_baseline(
        self, tmp_path: Path
    ) -> None:
        pdf = build_pdf(
            tmp_path / "a.pdf",
            [[Line("Uno")], [Line("[PO-SGS-400]", right="Política de Seguridad")]],
        )

        body = [b for b in read_pdf(pdf).blocks if not b.is_furniture]

        assert [(b.page, b.text) for b in body] == [
            (1, "Uno"),
            (2, "[PO-SGS-400] Política de Seguridad"),
        ]

    def test_repeated_header_and_footer_marked_as_furniture(self, tmp_path: Path) -> None:
        content = read_pdf(build_pdf(tmp_path / "m.pdf", manual_pages(), header=MANUAL_HEADER))

        furniture = {b.text for b in content.blocks if b.is_furniture}

        assert "Pza. Independencia 822 Piso 4" in furniture
        assert any(text.startswith("MC-SGG-435 / 23") for text in furniture)
        assert not any("REFERENCIAS" in text for text in furniture)

    def test_rotated_text_is_discarded(self, tmp_path: Path) -> None:
        content = read_pdf(build_pdf(tmp_path / "m.pdf", manual_pages(), header=MANUAL_HEADER))

        assert all(b.text != "CONFIDENCIAL TILSOR" for b in content.blocks)

    def test_bold_and_font_size(self, tmp_path: Path) -> None:
        content = read_pdf(
            build_pdf(tmp_path / "b.pdf", [[Line("4 REFERENCIAS", bold=True, size=13)]])
        )

        [block] = [b for b in content.blocks if b.text == "4 REFERENCIAS"]
        assert block.is_bold
        assert block.font_size == pytest.approx(13)

    def test_bookmarks_and_page_stats(self, tmp_path: Path) -> None:
        pdf = build_pdf(
            tmp_path / "s.pdf",
            [[Line("Texto suficiente para no parecer escaneado")], []],
            bookmarks=[[1, "4 REFERENCIAS", 1]],
            scanned_pages={2},
        )

        content = read_pdf(pdf)

        assert [(b.level, b.title, b.page) for b in content.bookmarks] == [(1, "4 REFERENCIAS", 1)]
        assert content.pages[1].image_coverage > 0.9
        assert content.pages[1].char_count == 0

    def test_encrypted_pdf(self, tmp_path: Path) -> None:
        pdf = build_pdf(tmp_path / "c.pdf", [[Line("secreto")]], encrypt=True)

        with pytest.raises(EncryptedDocumentError):
            read_pdf(pdf)

    def test_corrupt_pdf(self, tmp_path: Path) -> None:
        corrupt = tmp_path / "roto.pdf"
        corrupt.write_bytes(b"esto no es un pdf")

        with pytest.raises(ReadError):
            read_pdf(corrupt)


class TestDocxReader:
    def test_paragraphs_and_tables_interleaved_in_document_order(self, tmp_path: Path) -> None:
        path = build_docx(
            tmp_path / "o.docx",
            [Para("antes"), TableBlock([["a1", "a2"], ["b1", "b2"]]), Para("después")],
        )

        blocks = read_docx(path).blocks

        assert [b.text for b in blocks] == ["antes", "a1", "a2", "b1", "b2", "después"]
        assert [(b.table_index, b.row_index, b.cell_index) for b in blocks[1:5]] == [
            (0, 0, 0),
            (0, 0, 1),
            (0, 1, 0),
            (0, 1, 1),
        ]

    def test_merged_cells_are_read_once(self, tmp_path: Path) -> None:
        import docx

        path = tmp_path / "merge.docx"
        document = docx.Document()
        table = document.add_table(rows=1, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "combinada"
        table.cell(0, 2).text = "simple"
        document.save(str(path))

        assert [b.text for b in read_docx(path).blocks] == ["combinada", "simple"]

    def test_heading_level_from_internal_style_name_even_if_id_is_localized(
        self, tmp_path: Path
    ) -> None:
        path = build_docx(
            tmp_path / "h.docx",
            [Heading("REFERENCIAS", 1), Heading("Procedimientos", 2), Para("[PO-SGS-400]")],
            localized_heading_ids=True,
        )

        blocks = read_docx(path).blocks

        assert [b.heading_level for b in blocks] == [1, 2, None]

    def test_toc_style_is_flagged(self, tmp_path: Path) -> None:
        path = build_docx(tmp_path / "t.docx", [TocEntry("REFERENCIAS\t3"), Para("cuerpo")])

        assert [b.is_toc for b in read_docx(path).blocks] == [True, False]

    def test_textbox_read_once_and_kept_out_of_the_body(self, tmp_path: Path) -> None:
        path = build_docx(tmp_path / "x.docx", [TextBox("RG-SGSA-930 - Informe"), Para("cuerpo")])

        content = read_docx(path)

        assert content.textbox_texts == ["RG-SGSA-930 - Informe"]
        assert [b.text for b in content.blocks] == ["cuerpo"]

    def test_custom_properties_and_header(self, tmp_path: Path) -> None:
        path = build_docx(
            tmp_path / "p.docx",
            [Para("cuerpo")],
            custom_properties={"Referencia": "RG-SGSA-930", "Version": "01"},
            header_text="RG-SGSA-930 / 01 - Informe - Página 1/3",
        )

        content = read_docx(path)

        assert content.metadata["custom_properties"] == {
            "Referencia": "RG-SGSA-930",
            "Version": "01",
        }
        assert content.header_texts == ["RG-SGSA-930 / 01 - Informe - Página 1/3"]

    def test_corrupt_docx(self, tmp_path: Path) -> None:
        corrupt = tmp_path / "roto.docx"
        corrupt.write_bytes(b"no es un zip")

        with pytest.raises(ReadError):
            read_docx(corrupt)


class TestXlsxReader:
    def test_identity_only_with_properties(self, tmp_path: Path) -> None:
        sheet = build_xlsx(
            tmp_path / "RG-SGSA-927-01 Inventario.xlsx",
            title="Inventario de activos",
            custom_properties={"Referencia": "RG-SGSA-927", "Versión": "01"},
        )

        content = read_document(sheet)

        assert content.file_format is FileFormat.XLSX
        assert content.blocks == []  # las celdas no se analizan, aunque tengan códigos
        assert content.metadata["title"] == "Inventario de activos"
        assert content.metadata["sheets"] == ["Sheet"]
        assert content.metadata["custom_properties"] == {
            "Referencia": "RG-SGSA-927",
            "Versión": "01",
        }

    def test_password_protected_xlsx(self, tmp_path: Path) -> None:
        protected = tmp_path / "protegida.xlsx"
        protected.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 512)

        with pytest.raises(EncryptedDocumentError):
            read_xlsx(protected)

    def test_corrupt_xlsx(self, tmp_path: Path) -> None:
        corrupt = tmp_path / "rota.xlsx"
        corrupt.write_bytes(b"no es un zip")

        with pytest.raises(ReadError):
            read_xlsx(corrupt)


def test_unsupported_format(tmp_path: Path) -> None:
    other = tmp_path / "plantilla.dotx"
    other.write_bytes(b"")

    with pytest.raises(ReadError):
        read_document(other)
